"""Guidance, RAG assistant, bilingual UI, smart import and regulatory monitoring."""
import io
import zipfile

from docx import Document

from app import ai_review, catalogue, config, guidance, importer, monitor, rag
from app.db import Evidence, ImportFeedback, Project, RegAck, RegUpdate, Submission, User
from tests.conftest import login


def _sub(p, item_id, unit_code=None):
    return next(s for s in p.submissions if s.item_id == item_id and
                ((s.unit.code if s.unit else None) == unit_code))


# ---------------------------------------------------------------------------
# Knowledge & guidance
# ---------------------------------------------------------------------------

def test_every_item_has_a_cited_legal_level_and_guidance():
    assert len(catalogue.items()) == 73
    for it in catalogue.items():
        level, why, why_zh = guidance.level_of(it["id"])
        assert level in ("required", "supporting", "practice")
        assert why and why_zh, it["id"]
        g = guidance.item_guidance(it["id"])
        assert g["summary"] and g["steps"] and g["formats"]
        for m in g["minimum"]:
            assert len(m) == 3
    for e in rag.pool_entries():
        assert e["text"] and e["text_zh"] and e["kind"] in rag.pool()["kinds"]
        if e["src"]:
            assert e["src"] in rag.pool()["sources"]


def test_obligation_timeline_distinguishes_now_later_recommended(session):
    p = session.query(Project).filter_by(code="PRJ-2026-001").one()
    u = p.units[0]
    assert guidance.obligation(p, u, "C04-16").code == "mandatory_now"
    later = guidance.obligation(p, u, "C09-29")
    assert later.code == "mandatory_from" and later.applies_from.year == 2030
    assert guidance.obligation(p, u, "C11-40").code == "mandatory_now"      # Art 15(6) despite Art 12 being later
    assert guidance.obligation(p, u, "C04-73").code == "practice"
    assert guidance.obligation(p, u, "C02-07").code == "supporting"
    if not p.food_contact:
        assert guidance.obligation(p, u, "C04-17").code == "not_required"


def test_evaluate_finds_gaps_and_clears_them(session):
    p = session.query(Project).filter_by(code="PRJ-2026-002").one()
    s = _sub(p, "C03-10", "PU-001")
    ev = guidance.evaluate(p, s, data_override={"docno": "DW-1"}, evidence=[])
    assert any("revision" in g["en"].lower() or "uploaded" in g["en"].lower() for g in ev["gaps"])
    ok = guidance.evaluate(p, s, data_override={"docno": "DW-2026-010 R2"}, evidence=[])
    pattern = next(m for m in ok["minimum"] if m["spec"] == "pattern:docno")
    assert pattern["ok"] is True


def test_ai_checks_are_bilingual(session):
    p = session.query(Project).filter_by(code="PRJ-2026-001").one()
    s = _sub(p, "C04-16", "PU-001")
    _, checks = ai_review.review(s, p, s.unit, [])
    assert all("desc_zh" in c for c in checks)
    assert any(c["desc_zh"] for c in checks)


def test_rag_cross_lingual_and_citations():
    hits = rag.search("重金属限值是多少", corpus=False)
    assert hits and hits[0].id == "REG-A5-4"
    hits = rag.search("How long must the technical documentation be kept?", corpus=False)
    assert hits[0].id == "REG-A15-2"
    assert "STD-PRESUMPTION" in [h.id for h in rag.search("presumption of conformity EN 13430", corpus=False)[:2]]


def test_template_docx(client, session):
    login(client, "lumen")
    p = session.query(Project).filter_by(code="PRJ-2026-001").one()
    s = _sub(p, "C03-10", "PU-001")
    r = client.get(f"/projects/{p.id}/items/{s.id}/template.docx")
    assert r.status_code == 200
    text = "\n".join(par.text for par in Document(io.BytesIO(r.content)).paragraphs)
    assert "Annex VII" in text and "C03-10" in text


# ---------------------------------------------------------------------------
# Pages, language, pre-check, assistant
# ---------------------------------------------------------------------------

def test_new_pages_render_in_all_languages(client, session):
    login(client, "expert")
    p = session.query(Project).filter_by(code="PRJ-2026-002").one()
    urls = ["/updates", "/settings/ai", "/knowledge", f"/projects/{p.id}/guide", f"/projects/{p.id}/import",
            f"/projects/{p.id}/guide?level=mandatory_now", f"/updates/{session.query(RegUpdate).first().id}"]
    urls += [f"/projects/{p.id}/items/{s.id}" for s in p.submissions]
    for lang in ("bi", "en", "zh"):
        client.get(f"/lang/{lang}")
        for url in (urls if lang == "bi" else urls[:8]):
            r = client.get(url)
            assert r.status_code == 200, (lang, url)
    client.get("/lang/zh")
    assert "法规动态" in client.get("/").text
    client.get("/lang/en")
    assert "法规动态" not in client.get("/").text


def test_client_cannot_open_ai_settings(client):
    login(client, "lumen")
    r = client.get("/settings/ai", follow_redirects=False)
    assert r.status_code in (303, 403)


def test_precheck_is_a_dry_run(client, session):
    login(client, "aurora")
    p = session.query(Project).filter_by(code="PRJ-2026-002").one()
    s = _sub(p, "C03-15", "PU-001")
    before = s.status
    r = client.post(f"/projects/{p.id}/items/{s.id}/precheck", data={"f_rationale": "short"})
    body = r.json()
    assert body["ok"] and body["gaps"] and body["obligation"]["code"] == "mandatory_now"
    r = client.post(f"/projects/{p.id}/items/{s.id}/precheck", data={
        "f_rationale": "Protection: board grade needed to protect the product in transport. Manufacturing: fits the "
                       "existing line. Logistics: passes stacking. Functionality: reclosable lid."})
    assert not any("rationale" in g["en"].lower() and "100" in g["en"] for g in r.json()["gaps"])
    session.expire_all()
    assert session.get(Submission, s.id).status == before


def test_assistant_answers_with_sources_or_declines(client, session):
    login(client, "lumen")
    r = client.post("/assistant/ask", data={"q": "What are the PFAS limits for food-contact packaging?"}).json()
    assert r["ok"] and r["sources"] and "[S" in r["en"] and r["zh"]
    assert any(s["binding"] for s in r["sources"])
    r = client.post("/assistant/ask", data={"q": "zzqx wobble frobnicate"}).json()
    assert not r["sources"] and "TÜV SÜD expert" in r["en"]
    p = session.query(Project).filter_by(code="PRJ-2026-001").one()
    s = _sub(p, "C03-10", "PU-001")
    r = client.post("/assistant/ask", data={"q": "What do I need to provide?", "pid": p.id, "sid": s.id}).json()
    assert "C03-10" in r["en"]


# ---------------------------------------------------------------------------
# Smart import
# ---------------------------------------------------------------------------

def _docx(text):
    d = Document()
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_identification_rules():
    a = importer.analyze_name("PU-002_C04-16_report_R2_2026-03-01.pdf", "", [])
    assert a["item_id"] == "C04-16" and a["conf"] >= 0.9 and a["unit_ref"] == "PU-002" and a["date"] == "2026-03-01"
    b = importer.analyze_name("drawing.pdf", "03_Design_and_Manufacturing", [])
    assert b["item_id"].startswith("C03")
    cands = importer.content_candidates("Heavy metal test report. Lead < 5 mg/kg, cadmium < 2 mg/kg. Certificate no. 12")
    importer.lock_by_content(cands, "Heavy metal test report. Lead < 5 mg/kg, cadmium < 2 mg/kg. Certificate no. 12")
    assert any(c["id"] == "C04-70" and c["locked"] for c in cands)


def test_smart_import_end_to_end(client, session):
    login(client, "lumen")
    p = session.query(Project).filter_by(code="PRJ-2026-001").one()
    n_ev = len(p.evidence)
    z = io.BytesIO()
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("05_Article_5_Substances/lab.docx", _docx("Heavy metal test report — lead, cadmium, mercury and "
                                                              "chromium (VI) below 50 mg/kg. Certificate HM-77."))
        zf.writestr("misc/notes.txt", "unrelated text")
    files = [("files", ("PU-001_C03-10_DW-2026-010_R2.pdf", b"%PDF-1.4 drawing", "application/pdf")),
             ("files", ("bundle.zip", z.getvalue(), "application/zip"))]
    r = client.post(f"/projects/{p.id}/import/files", files=files).json()
    assert r["ok"], r
    assert client.get(f"/projects/{p.id}/import").status_code == 200
    session.expire_all()
    from app.db import ImportBatch
    b = session.query(ImportBatch).filter_by(project_id=p.id, status="staged").one()
    by = {f.filename: f for f in b.files}
    assert by["PU-001_C03-10_DW-2026-010_R2.pdf"].item_id == "C03-10"
    assert by["lab.docx"].item_id in ("C04-16", "C04-70")
    # the client corrects the lab report to C04-70 and skips the notes
    form = {f"item_{by['lab.docx'].id}": "C04-70", f"item_{by['notes.txt'].id}": "skip",
            "default_unit": str(p.units[0].id), "auto_submit": "1", "auto_ai": "1"}
    r = client.post(f"/projects/{p.id}/import/confirm", data=form).json()
    assert r["ok"], r
    session.expire_all()
    p = session.get(Project, p.id)
    assert len(p.evidence) == n_ev + 2
    e = next(e for e in p.evidence if e.item_id == "C04-70")
    assert e.code.startswith("PU-001-")
    s = _sub(p, "C04-70", "PU-001")
    assert s.status in ("ai_reviewed", "approved")
    if by["lab.docx"].proposed_item != "C04-70":
        assert session.query(ImportFeedback).filter_by(item_id="C04-70").count() == 1


def test_table_import_creates_units_and_fields(client, session):
    login(client, "lumen")
    p = session.query(Project).filter_by(code="PRJ-2026-001").one()
    n_units = len(p.units)
    client.post(f"/projects/{p.id}/import/table", data={"text": importer.SAMPLE_TABLE})
    r = client.post(f"/projects/{p.id}/import/confirm", data={"auto_submit": "1"}).json()
    assert r["ok"], r
    session.expire_all()
    p = session.get(Project, p.id)
    assert len(p.units) == n_units + 2
    u = next(u for u in p.units if u.upi == "DEMO-CASE-48")
    assert u.level == "Transport" and u.mass_g == 480
    assert _sub(p, "C04-16", u.code).data.get("res", "").startswith("< 60")


def test_table_rows_without_name_fill_the_default_unit(client, session):
    login(client, "lumen")
    p = session.query(Project).filter_by(code="PRJ-2026-001").one()
    n_units, u = len(p.units), p.units[0]
    client.post(f"/projects/{p.id}/import/table", data={"text": "Supplier\t重金属\t检测方法\nAcme\t< 30\tICP-OES"})
    client.post(f"/projects/{p.id}/import/confirm", data={"default_unit": str(u.id)})
    session.expire_all()
    p = session.get(Project, p.id)
    assert len(p.units) == n_units


def test_import_onto_approved_item_reopens_it(client, session):
    login(client, "lumen")
    p = session.query(Project).filter_by(code="PRJ-2026-001").one()
    s = next(x for x in p.submissions if x.status == "approved" and x.item_id == "C03-10")
    code = s.unit.code
    client.post(f"/projects/{p.id}/import/files",
                files=[("files", (f"{code}_C03-10_drawing_R9.pdf", b"%PDF-1.4 x", "application/pdf"))])
    client.post(f"/projects/{p.id}/import/confirm", data={"default_unit": str(s.unit_id)})
    session.expire_all()
    assert session.get(Submission, s.id).status in ("submitted", "ai_reviewed")


def test_upload_warns_when_file_looks_like_another_item(client, session):
    login(client, "lumen")
    p = session.query(Project).filter_by(code="PRJ-2026-001").one()
    s = _sub(p, "C03-10", "PU-001")
    r = client.post(f"/projects/{p.id}/evidence", data={"sub_id": s.id, "type_code": "DRW"},
                    files={"file": ("C04-16_heavy_metal_report.pdf", b"%PDF-1.4 x", "application/pdf")},
                    follow_redirects=False)
    assert "toast_type=warn" in r.headers["location"]


# ---------------------------------------------------------------------------
# Monitoring
# ---------------------------------------------------------------------------

def test_updates_seeded_with_impact_and_ack(client, session):
    assert session.query(RegUpdate).count() >= 8
    login(client, "lumen")
    user = session.query(User).filter_by(username="lumen").one()
    before = monitor.unread_count(session, user)
    u = session.query(RegUpdate).filter_by(code="REG-2026-001").one()
    assert client.get(f"/updates/{u.id}").status_code == 200
    session.expire_all()
    assert monitor.unread_count(session, user) == before - 1
    client.post("/updates/ack-all")
    session.expire_all()
    assert monitor.unread_count(session, user) == 0
    projects = session.query(Project).filter_by(client_id=user.client_id).all()
    assert monitor.impact(u, projects)


def test_scanner_detects_new_knowledge_documents(session, tmp_path, monkeypatch):
    kb = tmp_path / "kb"
    (kb / "03 consultation").mkdir(parents=True)
    monkeypatch.setattr(config, "KNOWLEDGE_DIR", kb)
    monkeypatch.setattr(monitor, "MANIFEST_PATH", tmp_path / "manifest.json")
    assert monitor.scan(session)["baseline"]
    (kb / "03 consultation" / "draft_act.html").write_text(
        "<html><body><h1>Draft implementing act</h1><p>Methodology for recycled content in plastic packaging "
        "and verification of post-consumer recyclate.</p></body></html>", encoding="utf-8")
    res = monitor.scan(session)
    session.flush()
    assert len(res["created"]) == 1
    u = session.query(RegUpdate).filter_by(code=res["created"][0]).one()
    assert "a7" in u.articles and "C06-22" in u.items and not u.reviewed


def test_expert_recheck_reopens_approved_items(client, session):
    login(client, "expert")
    p = session.query(Project).filter_by(code="PRJ-2026-001").one()
    u = session.query(RegUpdate).filter_by(code="REG-2026-001").one()
    approved = [s for s in p.submissions if s.item_id in u.items and s.status == "approved"]
    client.post(f"/updates/{u.id}/recheck", data={"project_id": p.id})
    session.expire_all()
    for s in approved:
        assert session.get(Submission, s.id).status == "submitted"
