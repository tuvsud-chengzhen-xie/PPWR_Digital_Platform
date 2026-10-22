"""Smart import — the V4.2 prototype's one-click upload + AI identification, server side.

Files (single, multiple, a whole folder or a ZIP) are identified against the 73 catalogue items:

  A. item code in the name        C04-16, c04_16, C0416                    → 95 %
  B. category-item number          04-16, 04_16 (validated, not a date)     → 90 %
  C. keyword library (中/EN)      per item ≥ 8 terms, folder hint narrows  → ≤ 88 %
  D. packaging unit               PU-002, PKG-02 or a registered UPI in the path
  E. version / date               R2, Rev 3, v1.2 · 2026-08-30
  F. content                      PDF text (pypdf), DOCX, XLSX, TXT/CSV — keyword scoring plus the
                                  V4.2 lock rules (e.g. heavy-metal terms + "certificate" → C04-70)
  G. feedback learning            a human correction is remembered per client and re-weights the next
                                  identification of a file with the same base name

Everything is staged first (ImportBatch / ImportFile); the client reviews the proposals, corrects
item / unit / type where needed, and imports with one click. Import creates evidence with proper
IDs, optionally submits the items and runs the AI pre-review — the same workflow as manual upload.

Tables (pasted TSV/CSV or an XLSX/CSV file) are parsed with the V4.2 header-alias library into
packaging-unit data, BOM lines and item fields.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
import shutil
import zipfile
from functools import lru_cache
from pathlib import Path

from sqlalchemy.orm import Session

from . import ai_review, catalogue, config, storage, workflow
from .db import (Component, Evidence, ImportBatch, ImportFeedback, ImportFile, PackagingUnit, Project,
                 Submission, Supplier, log, now)

STAGING = config.UPLOAD_DIR / "_staging"
MAX_FILES = 300
MAX_TOTAL_MB = 200
TEXT_EXT = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".htm", ".log"}
LEVEL_MAP = [(r"销售|sales|retail", "Sales"), (r"组合|grouped|group|secondary", "Grouped"),
             (r"运输|transport|shipping|tertiary|pallet", "Transport"), (r"电商|e-?commerce", "E-commerce")]


@lru_cache(maxsize=1)
def rules() -> dict:
    return json.loads((config.SEEDS_DIR / "import_rules.json").read_text(encoding="utf-8"))


# ---------------------------------------------------------------------------
# Identification
# ---------------------------------------------------------------------------

def _kw_match(s: str, cat_filter: int = 0):
    best = None
    for iid, kws in rules()["keywords"].items():
        if cat_filter and int(iid[1:3]) != cat_filter:
            continue
        sc, hits = 0.0, []
        for k in kws:
            kl = k.lower()
            if kl and kl in s:
                sc += min(len(kl), 8) / 8
                hits.append(k)
        if sc > 0 and (best is None or sc > best["sc"]):
            best = {"id": iid, "sc": sc, "hits": hits}
    return best


def folder_cat(folder: str) -> int:
    f = (folder or "").lower()
    for x in rules()["folders"]:
        if x["folder"].lower() in f and x["cats"]:
            return x["cats"][0]
    return 0


def analyze_name(fname: str, folder: str, units: list[PackagingUnit]) -> dict:
    path = (folder + "/" if folder else "") + fname
    s = path.lower()
    out = {"item_id": None, "conf": 0.0, "rule": [], "unit_ref": "", "unit_id": None, "version": "", "date": ""}
    m = re.search(r"c(\d{2})[-_. ]?(\d{2})(?![0-9])", s)
    if m and catalogue.item(f"C{m.group(1)}-{m.group(2)}"):
        out.update(item_id=f"C{m.group(1)}-{m.group(2)}", conf=0.95)
        out["rule"].append(f"code {out['item_id']} · 编号")
    if not out["item_id"]:
        m = re.search(r"(?:^|[^0-9.])(\d{2})[-_.](\d{2})(?![0-9])", s)
        if m:
            cat, no = int(m.group(1)), int(m.group(2))
            cands = [i for i in catalogue.items() if i["cat"] == cat and int(i["id"][4:]) == no] if 1 <= cat <= 16 else []
            if len(cands) == 1:
                out.update(item_id=cands[0]["id"], conf=0.9)
                out["rule"].append(f"number {m.group(1)}-{m.group(2)} · 编号")
    if not out["item_id"]:
        fc = folder_cat(folder)
        best = _kw_match(s, fc) or (_kw_match(s, 0) if fc else None)
        if best:
            bonus = 0.05 if fc and int(best["id"][1:3]) == fc else 0
            out.update(item_id=best["id"], conf=min(0.88, 0.45 + best["sc"] * 0.18 + bonus))
            out["rule"].append("keywords 关键词: " + " / ".join(best["hits"][:3]))
    m = re.search(r"(?<![A-Za-z])PU[-_ ]?(\d{1,3})(?!\d)", path, re.I) or re.search(r"PKG[-_ ]?(\d{1,3})(?!\d)", path, re.I)
    if m:
        out["unit_ref"] = f"PU-{int(m.group(1)):03d}"
    else:
        for u in units:
            if u.upi and len(u.upi) > 3 and u.upi.upper() in path.upper():
                out["unit_ref"] = u.code
                break
    if out["unit_ref"]:
        u = next((u for u in units if u.code == out["unit_ref"]), None)
        out["unit_id"] = u.id if u else None
        out["rule"].append(f"unit 包装 {out['unit_ref']}" + ("" if u else " (not registered 未登记)"))
    m = re.search(r"(?:^|[_\-\s])(?:r|rev|v|ver|version)[.\s]?(\d{1,2}(?:\.\d{1,2})?)(?![\d])", path, re.I)
    if m:
        out["version"] = m.group(0).strip(" _-")
    m = re.search(r"(20\d{2})[-_./年](\d{1,2})[-_./月](\d{1,2})", path)
    if m:
        out["date"] = f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    return out


def extract_text(path: Path) -> tuple[str, str]:
    """(text, method). Never raises — unreadable files simply have no content signal."""
    ext = path.suffix.lower()
    try:
        if ext in TEXT_EXT:
            raw = path.read_bytes()[:400_000]
            for enc in ("utf-8-sig", "gb18030", "latin-1"):
                try:
                    return raw.decode(enc), "text"
                except UnicodeDecodeError:
                    continue
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            parts = []
            for pg in reader.pages[:8]:
                parts.append(pg.extract_text() or "")
                if sum(map(len, parts)) > 80_000:
                    break
            text = "\n".join(parts)
            return text, "pdf-text" if text.strip() else "pdf (no text layer — scanned?)"
        if ext == ".docx":
            from docx import Document
            d = Document(str(path))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for r in t.rows:
                    parts.append(" | ".join(c.text for c in r.cells))
            return "\n".join(parts), "docx"
        if ext in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            wb = load_workbook(str(path), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets[:5]:
                for row in ws.iter_rows(max_row=200, values_only=True):
                    parts.append(" ".join(str(c) for c in row if c is not None))
            return "\n".join(parts), "xlsx"
        if ext in (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".gif"):
            return "", "image (no OCR in this pilot)"
    except Exception as exc:  # corrupt / protected files
        return "", f"unreadable ({type(exc).__name__})"
    return "", "filename only"


def content_candidates(text: str) -> list[dict]:
    s = (text or "").lower()
    out = []
    for iid, kws in rules()["keywords"].items():
        sc, hits, lines = 0.0, [], []
        for k in kws:
            kl = k.lower()
            idx = s.find(kl) if kl else -1
            if idx >= 0:
                sc += min(len(kl), 8) / 8
                hits.append(k)
                if len(lines) < 3:
                    a = s.rfind("\n", 0, idx) + 1
                    b = s.find("\n", idx)
                    line = text[a:b if b >= 0 else len(s)].strip()
                    if 5 < len(line) < 200:
                        lines.append(line)
        if sc > 0:
            out.append({"id": iid, "sc": sc, "hits": hits, "lines": lines, "locked": False})
    out.sort(key=lambda c: -c["sc"])
    return out


def lock_by_content(cands: list[dict], text: str) -> dict:
    """The V4.2 lock rules: strong topical evidence pins the item."""
    s = (text or "").lower()
    n = lambda rx: len(re.findall(rx, s))  # noqa: E731
    hm = n(r"重金属|铅|镉|汞|六价铬|heavy metals?|cadmium|\blead\b|mercury|chromium ?\(?vi\)?|cr6\+")
    pf = n(r"pfas|全氟|多氟|perfluoro|pfoa|pfos")
    fcm = n(r"食品接触|food contact|\bfcm\b|1935/2004")
    recl = n(r"可回收|recyclab|回收性")
    coc = n(r"\bcoc\b|符合性声明|符合性证书|declaration of conformity|statement of conformity|certificate|证书|test report|检测报告")
    tf = n(r"total fluorine|总氟|燃烧离子色谱|\bcic\b")
    a5 = n(r"article\s*5\b|art\.?\s*5\b|substances? of concern")

    def lock(iid, strength):
        """strength = how specific the evidence is; the most specific lock wins (V4.2 rules + weighting)."""
        c = next((c for c in cands if c["id"] == iid), None)
        if c is None:
            c = {"id": iid, "sc": 0.5, "hits": [], "lines": [], "locked": False}
            cands.append(c)
        c["locked"], c["lock"] = True, max(strength, c.get("lock", 0))
    if hm >= 2:
        lock("C04-16", 2)
    if pf >= 2:
        lock("C04-17", 2)
    if fcm >= 2:
        lock("C04-18", 1)
    if recl >= 2:
        lock("C05-19", 2)
    if hm >= 1 and coc >= 1:
        lock("C04-70", 3)
    if pf >= 1 and (tf >= 1 or coc >= 1):
        lock("C04-71", 3 + (1 if tf else 0))
    if fcm >= 1 and coc >= 1 and pf == 0:
        lock("C04-72", 3)
    if a5 >= 1 and coc >= 1 and (hm + pf) >= 2:
        lock("C04-73", 2)
    limit = n(r"[<≤]\s*\d+(?:\.\d+)?\s*(?:mg/kg|ppm|ppb|%)")
    return {"hm": hm, "pfas": pf, "fcm": fcm, "recyclability": recl, "certificate": coc, "limit": limit}


def _stem(name: str) -> str:
    b = re.sub(r"\.[a-z0-9]{2,5}$", "", name.lower())
    b = re.sub(r"[_\-\s]?(v\d|r\d|rev\d|ver\d|final|草稿|定稿|\(\d+\)|\d{4}-\d{2}-\d{2}).*$", "", b)
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]+", " ", b).strip()[:60]


def feedback_hits(session: Session, client_id: int | None, name: str) -> dict[str, float]:
    stem = _stem(name)
    if not stem:
        return {}
    hits: dict[str, float] = {}
    for fb in session.query(ImportFeedback).filter(ImportFeedback.client_id == client_id):
        if fb.stem == stem:
            hits[fb.item_id] = hits.get(fb.item_id, 0) + min(fb.count, 3) * 0.05
        elif len(stem) > 4 and len(fb.stem) > 4 and (fb.stem in stem or stem in fb.stem):
            hits[fb.item_id] = hits.get(fb.item_id, 0) + min(fb.count, 2) * 0.025
    return hits


def record_feedback(session: Session, client_id: int | None, name: str, item_id: str) -> None:
    stem = _stem(name)
    if not stem or not item_id:
        return
    fb = session.query(ImportFeedback).filter_by(client_id=client_id, stem=stem, item_id=item_id).first()
    if fb:
        fb.count += 1
    else:
        session.add(ImportFeedback(client_id=client_id, stem=stem, item_id=item_id))


def guess_type(item_id: str | None, ext: str, text: str) -> str:
    s = (text or "").lower()
    if ext in (".dwg", ".dxf"):
        return "DRW"
    if ext in (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp", ".gif"):
        return "PHO" if not item_id or catalogue.item(item_id)["evidence_code"] != "DRW" else "DRW"
    if re.search(r"test report|检测报告|laboratory|实验室|analysis report", s):
        return "TST"
    if re.search(r"certificate|证书|\bcoc\b", s):
        return "CRT" if not item_id or catalogue.item(item_id)["evidence_code"] not in ("TST",) else "TST"
    if re.search(r"declaration|声明|statement", s):
        return "DCL"
    return catalogue.item(item_id)["evidence_code"] if item_id and catalogue.item(item_id) else "DOC"


def identify(session: Session, project: Project, path: Path, fname: str, folder: str) -> dict:
    out = analyze_name(fname, folder, project.units)
    text, method = extract_text(path)
    out["method"], out["snippet"], out["cands"] = method, "", []
    if text.strip():
        cands = content_candidates(text)
        locks = lock_by_content(cands, text)
        out["cands"] = [{"id": c["id"], "sc": round(c["sc"], 2), "hits": c["hits"][:4], "locked": c["locked"]}
                        for c in cands[:5]]
        snippet = []
        for c in cands[:3]:
            snippet += [ln for ln in c["lines"] if ln not in snippet]
        out["snippet"] = " | ".join(snippet[:4])[:900]
        locked = [c for c in cands if c["locked"]]
        top = max(locked, key=lambda c: (c.get("lock", 0), c["sc"])) if locked else (cands[0] if cands else None)
        if top:
            bonus = (0.05 if locks["limit"] else 0) + (0.03 if re.search(r"供应商|supplier|vendor|manufacturer", text, re.I) else 0)
            name_sc = out["conf"]
            blended = min(0.96, name_sc * 0.4 + min(top["sc"], 3) * 0.18 + bonus)
            if name_sc >= 0.9 and out["item_id"] != top["id"] and not top["locked"]:
                pass  # an explicit item code in the file name wins over unlocked content
            else:
                if top.get("locked"):
                    blended = max(blended, min(0.92, 0.72 + 0.05 * top.get("lock", 1)))
                out["item_id"], out["conf"] = top["id"], max(blended, name_sc * 0.9 if out["item_id"] == top["id"] else blended)
                out["rule"].append("content 内容: " + " / ".join(top["hits"][:3]) + (" (locked 锁定)" if top["locked"] else ""))
        if not out["date"]:
            m = re.search(r"(20\d{2})[-./年](\d{1,2})[-./月](\d{1,2})", text)
            if m:
                try:
                    out["date"] = dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
                except ValueError:
                    pass
    fb = feedback_hits(session, project.client_id, fname)
    if fb:
        best_id, best_v = max(fb.items(), key=lambda x: x[1])
        if best_id != out["item_id"] and best_v >= 0.1:
            out["item_id"], out["conf"] = best_id, min(0.95, (out["conf"] or 0.5) + best_v)
            out["rule"].append(f"learned from corrections 反馈学习 → {best_id}")
        elif best_id == out["item_id"]:
            out["conf"] = min(0.96, out["conf"] + best_v)
    out["type_code"] = guess_type(out["item_id"], path.suffix.lower(), text)
    return out


# ---------------------------------------------------------------------------
# Staging
# ---------------------------------------------------------------------------

def open_batch(session: Session, project: Project, actor: str) -> ImportBatch:
    b = session.query(ImportBatch).filter_by(project_id=project.id, status="staged").first()
    if b is None:
        b = ImportBatch(project_id=project.id, created_by=actor)
        session.add(b)
        session.flush()
    return b


def _expand(files: list[tuple[str, bytes]]) -> list[tuple[str, str, bytes]]:
    """(name, data) uploads → (folder, filename, data), unpacking ZIPs."""
    out, total = [], 0
    for name, data in files:
        name = (name or "file").replace("\\", "/")
        folder, _, fname = name.rpartition("/")
        if fname.lower().endswith(".zip"):
            try:
                with zipfile.ZipFile(io.BytesIO(data)) as z:
                    for info in z.infolist():
                        if info.is_dir() or "__MACOSX" in info.filename or info.filename.rsplit("/", 1)[-1].startswith("."):
                            continue
                        total += info.file_size
                        if total > MAX_TOTAL_MB * 1024 * 1024 or len(out) >= MAX_FILES:
                            break
                        inner = info.filename
                        try:
                            inner = inner.encode("cp437").decode("utf-8")
                        except (UnicodeEncodeError, UnicodeDecodeError):
                            pass
                        f2, _, n2 = inner.rpartition("/")
                        out.append((f2 or fname[:-4], n2, z.read(info)))
            except zipfile.BadZipFile:
                out.append((folder, fname, data))
            continue
        total += len(data)
        if total > MAX_TOTAL_MB * 1024 * 1024 or len(out) >= MAX_FILES:
            break
        out.append((folder, fname, data))
    return out


def stage_files(session: Session, project: Project, files: list[tuple[str, bytes]], actor: str) -> dict:
    batch = open_batch(session, project, actor)
    d = STAGING / str(batch.id)
    d.mkdir(parents=True, exist_ok=True)
    added = []
    for folder, fname, data in _expand(files):
        if not data:
            continue
        safe = storage.safe_name(fname)
        path = d / f"{len(batch.files) + len(added) + 1:04d}__{safe}"
        path.write_bytes(data)
        r = identify(session, project, path, fname, folder)
        f = ImportFile(batch_id=batch.id, filename=fname, folder=folder, stored_path=str(path), size=len(data),
                       item_id=r["item_id"], proposed_item=r["item_id"], unit_id=r["unit_id"], unit_ref=r["unit_ref"],
                       type_code=r["type_code"], confidence=round(r["conf"], 2), rule=" · ".join(r["rule"]),
                       method=r["method"], snippet=r["snippet"], version=r["version"], found_date=r["date"],
                       candidates=r["cands"])
        session.add(f)
        added.append(f)
    session.flush()
    hi = sum(1 for f in added if f.item_id and f.confidence >= 0.9)
    mid = sum(1 for f in added if f.item_id and 0.7 <= f.confidence < 0.9)
    lo = sum(1 for f in added if f.item_id and f.confidence < 0.7)
    none = sum(1 for f in added if not f.item_id)
    return {"batch": batch, "added": len(added), "hi": hi, "mid": mid, "lo": lo, "none": none}


def update_files(session: Session, batch: ImportBatch, form) -> int:
    n = 0
    for f in batch.files:
        if f.status != "staged":
            continue
        item = form.get(f"item_{f.id}")
        if item is not None:
            item = item if item and item != "skip" and catalogue.item(item) else None
            if item != f.item_id:
                f.item_id, n = item, n + 1
        unit = form.get(f"unit_{f.id}")
        if unit is not None:
            f.unit_id = int(unit) if unit.isdigit() else None
        typ = form.get(f"type_{f.id}")
        if typ and typ in catalogue.EVIDENCE_TYPES:
            f.type_code = typ
    return n


# ---------------------------------------------------------------------------
# Tables
# ---------------------------------------------------------------------------

def parse_delimited(text: str) -> list[list[str]]:
    lines = [ln for ln in (text or "").replace("\r", "").lstrip("\ufeff").split("\n") if ln.strip()]
    if not lines:
        return []
    if any("\t" in ln for ln in lines[:3]):
        return [[c.strip() for c in ln.split("\t")] for ln in lines]
    return [[c.strip() for c in row] for row in csv.reader(lines)]


def parse_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    return [["" if c is None else str(c).strip() for c in row] for row in ws.iter_rows(max_row=500, values_only=True)]


def match_alias(cell: str, has_bom: bool) -> str:
    s = (cell or "").strip().lower()
    if not s:
        return ""
    best, best_w = "", 0
    mass_like = bool(re.match(r"^(质量|重量|克重|mass|weight)", s))
    sup_like = s in ("供应商", "supplier", "厂商", "vendor")
    for target, arr in rules()["aliases"].items():
        for a in arr:
            al = a.lower()
            exact = s == al
            incl = not exact and ((al in s and len(al) >= 2) or (s in al and len(s) >= 2))
            if not exact and not incl:
                continue
            if mass_like and target.startswith("pkg.") and has_bom and \
                    not re.match(r"^(包装质量|总质量|净重|单位质量|包装总质量|包装净重)$", s):
                continue
            if mass_like and target.startswith("bom.") and not has_bom:
                continue
            if sup_like and has_bom and target == "pkg.sup":
                continue
            if sup_like and not has_bom and target.startswith("bom."):
                continue
            w = (100 if exact else 0) + min(len(al), 8)
            if w > best_w:
                best, best_w = target, w
    return best


def targets() -> list[tuple[str, str]]:
    out = [("", "— ignore 忽略 —")]
    labels = {"pkg.name": "Unit · name 名称", "pkg.type": "Unit · level 层级", "pkg.upi": "Unit · UPI",
              "pkg.skus": "Unit · SKU / product 产品", "pkg.mass": "Unit · mass g 质量", "pkg.dims": "Unit · dimensions 尺寸",
              "pkg.sup": "Unit · supplier 供应商", "pkg.site": "Unit · site 工厂", "pkg.country": "Unit · country 国家",
              "pkg.ref": "Unit · photo / artwork ref", "bom.0": "BOM · component 组件", "bom.1": "BOM · material 材料",
              "bom.3": "BOM · mass g 质量", "bom.4": "BOM · dimensions 尺寸", "bom.5": "BOM · supplier 供应商",
              "remark": "Remark 备注"}
    for t in rules()["aliases"]:
        if t.startswith("f."):
            iid, k = t[2:].split(".")
            it = catalogue.item(iid)
            f = next((x for x in (it or {}).get("fields", []) if x["k"] == k), None)
            if f:
                labels[t] = f"{iid} · {f['label']}"
    return out + list(labels.items())


def build_table(grid: list[list[str]], col_map: list[str] | None = None) -> dict:
    grid = [r for r in grid if any(str(c).strip() for c in r)]
    if not grid:
        return {}
    head = 0
    best = 0
    for i, row in enumerate(grid[:5]):
        n = sum(1 for c in row if c and match_alias(str(c), False))
        if n > best:
            head, best = i, n
    header = [str(c) for c in grid[head]]
    has_bom = any(str(t).startswith("bom.") for t in col_map) if col_map else \
        any(match_alias(c, False) == "bom.0" for c in header)
    col_map = col_map or [match_alias(c, has_bom) for c in header]
    col_map = (col_map + [""] * len(header))[:len(header)]
    return {"grid": grid[:300], "head": head, "header": header, "map": col_map}


def table_rows(table: dict) -> list[dict]:
    rows = []
    for r in table.get("grid", [])[table.get("head", 0) + 1:]:
        row = {"pkg": {}, "bom": {}, "fields": {}, "remark": ""}
        for ci, target in enumerate(table.get("map", [])):
            v = str(r[ci]).strip() if ci < len(r) and r[ci] is not None else ""
            if not v or not target:
                continue
            if target == "remark":
                row["remark"] = v
            elif target.startswith("pkg."):
                row["pkg"][target[4:]] = v
            elif target.startswith("bom."):
                row["bom"][int(target[4:])] = v
            elif target.startswith("f."):
                iid, k = target[2:].split(".")
                row["fields"].setdefault(iid, {})[k] = v
        if row["pkg"] or row["bom"] or row["fields"]:
            rows.append(row)
    return rows


SAMPLE_TABLE = "\n".join([
    "包装名称 Packaging name\t包装类型 Type\tUPI\t产品/SKU\t质量(g)\t尺寸(mm)\t供应商 Supplier\t生产国 Country\t重金属检测结果\t检测方法\t可回收性结论\t备注",
    "Colour card A5 (import demo)\t销售包装 Sales\tDEMO-CARD-A5\tSKU-A / SKU-B\t14\t210×148×0.4\tHengda Paperboard Co., Ltd.\tChina\t< 50 mg/kg\tEN 13657 + ICP-OES\tMono-material paperboard, paper & board stream; EN 13430 used as guidance\tnon food-contact",
    "Shipper case (import demo)\t运输包装 Transport\tDEMO-CASE-48\tSKU-A / SKU-B\t480\t600×400×350\tRhein Wellpappe GmbH\tGermany\t< 60 mg/kg\tEN 13657 + ICP-OES\tCorrugated board, paper & board stream\tBOPP tape on the case",
])


def _num(v):
    m = re.search(r"(\d+(?:[.,]\d+)?)", str(v or ""))
    return float(m.group(1).replace(",", ".")) if m else None


def _level(v: str) -> str:
    for rx, lvl in LEVEL_MAP:
        if re.search(rx, v or "", re.I):
            return lvl
    return "Sales"


def _option(field: dict, v: str) -> str:
    opts, zh = field.get("options", []), field.get("options_zh", [])
    low = v.strip().lower()
    for i, o in enumerate(opts):
        if low == o.lower() or (i < len(zh) and v.strip() == zh[i]):
            return o
    if re.search(r"不适用|^否|^no\b|n/a|not|none|single", low):
        return next((o for o in opts if re.match(r"(Not|No|None|Single)", o)), v)
    if re.search(r"^是|^yes|applicable|适用|reusable|present", low):
        return next((o for o in opts if re.match(r"(Yes|Applicable|Reusable|Present)", o)), v)
    return v


def apply_table(session: Session, project: Project, table: dict, default_unit: PackagingUnit | None, actor: str,
                touched: set) -> dict:
    stats = {"rows": 0, "units_new": 0, "units_updated": 0, "bom": 0, "fields": 0, "remarks": 0}
    suppliers = {s.name.lower(): s for s in session.query(Supplier).filter_by(client_id=project.client_id)}
    for row in table_rows(table):
        stats["rows"] += 1
        unit = default_unit
        p = row["pkg"]
        ref = p.get("upi") or p.get("name")
        if p and not ref and default_unit is None:
            p = {}
        if p:
            # a row without name / UPI describes the default unit — never invent a new unit for it
            unit = next((u for u in project.units if ref and (u.upi == ref or u.name == ref or u.code == ref)),
                        None) if ref else default_unit
            if unit is None:
                unit = PackagingUnit(project=project, code=workflow.next_unit_code(project), name=p.get("name") or ref or "Imported unit")
                session.add(unit)
                session.flush()
                session.refresh(project)
                workflow.ensure_submissions(session, project)
                stats["units_new"] += 1
            else:
                stats["units_updated"] += 1
            if p.get("name"):
                unit.name = p["name"]
            if p.get("type"):
                unit.level = _level(p["type"])
            for k, attr in (("upi", "upi"), ("skus", "skus"), ("dims", "dims"), ("sup", "supplier_name"),
                            ("site", "site"), ("country", "country"), ("ref", "artwork_ref")):
                if p.get(k):
                    setattr(unit, attr, p[k])
            if p.get("mass") and _num(p["mass"]) is not None:
                unit.mass_g = _num(p["mass"])
            session.flush()
            workflow.sync_unit_register(session, unit, actor)
            for s in project.submissions:
                if s.unit_id == unit.id and s.item_id in workflow.UNIT_REGISTER_ITEMS:
                    touched.add(s.id)
        if unit is None and project.units:
            unit = project.units[0]
        if row["bom"] and unit is not None:
            b = row["bom"]
            c = Component(unit=unit, code=workflow.next_component_code(unit), pack_type=b.get(0, ""),
                          description=b.get(0, ""), material=b.get(1, ""),
                          material_group=catalogue.material_group_for(b.get(1, "")), mass_g=_num(b.get(3)),
                          dims=b.get(4, ""))
            sup = suppliers.get((b.get(5) or "").lower())
            if sup:
                c.supplier_id = sup.id
            elif b.get(5):
                c.remarks = f"Supplier: {b[5]}"
            if row["remark"]:
                c.remarks = (c.remarks + "; " if c.remarks else "") + row["remark"]
                stats["remarks"] += 1
            session.add(c)
            session.flush()
            stats["bom"] += 1
        for iid, values in row["fields"].items():
            it = catalogue.item(iid)
            if not it:
                continue
            uid = None if it["scope"] == "project" else (unit.id if unit else None)
            sub = next((s for s in project.submissions if s.item_id == iid and s.unit_id == uid), None)
            if sub is None or (sub.status == "approved"):
                continue
            data = dict(sub.data or {})
            for k, v in values.items():
                f = next((x for x in it.get("fields", []) if x["k"] == k), None)
                if f is None:
                    continue
                data[k] = _option(f, v) if f["type"] == "sel" else v
                stats["fields"] += 1
            sub.data = data
            touched.add(sub.id)
    session.flush()
    session.refresh(project)
    workflow.sync(session, project)
    return stats


# ---------------------------------------------------------------------------
# One-click import
# ---------------------------------------------------------------------------

def confirm(session: Session, project: Project, batch: ImportBatch, actor: str, auto_submit: bool,
            auto_ai: bool, default_unit_id: int | None) -> dict:
    default_unit = next((u for u in project.units if u.id == default_unit_id), project.units[0] if project.units else None)
    touched: set[int] = set()
    stats = {"files": 0, "skipped": 0, "corrections": 0, "approved_attached": 0, "submitted": 0, "ai": 0}
    for f in batch.files:
        if f.status != "staged":
            continue
        if not f.item_id:
            f.status = "skipped"
            stats["skipped"] += 1
            continue
        it = catalogue.item(f.item_id)
        unit = None if it["scope"] == "project" else next((u for u in project.units if u.id == f.unit_id), default_unit)
        uid = unit.id if unit else None
        sub = next((s for s in project.submissions if s.item_id == f.item_id and s.unit_id == uid), None)
        if sub is None or sub.status == "na":
            f.status = "skipped"
            stats["skipped"] += 1
            continue
        if f.proposed_item and f.proposed_item != f.item_id:
            record_feedback(session, project.client_id, f.filename, f.item_id)
            stats["corrections"] += 1
        data = Path(f.stored_path).read_bytes() if f.stored_path and Path(f.stored_path).exists() else b""
        code = workflow.next_evidence_code(session, project, unit, None, f.type_code)
        path = storage.save(project.code, code, storage.safe_name(f.filename), data)
        issued = None
        if f.found_date and f.type_code in ("TST", "CRT", "DCL"):
            try:
                issued = dt.date.fromisoformat(f.found_date)
            except ValueError:
                issued = None
        e = Evidence(project=project, unit_id=uid, item_id=f.item_id, code=code,
                     title=Path(f.filename).stem.replace("_", " ")[:120], type_code=f.type_code,
                     file_no=f.version, issued_date=issued, responsible="Applicant", conformity="Not stated",
                     filename=storage.safe_name(f.filename), stored_path=str(path), size=len(data), uploaded_by=actor)
        session.add(e)
        session.flush()
        f.status, f.evidence_id = "imported", e.id
        stats["files"] += 1
        if sub.status == "approved":
            # new evidence on an approved item must be reviewed before it can reach the TD
            sub.status, sub.reopened_reason = "submitted", f"New evidence {code} added by smart import — re-review."
            sub.ai_verdict, sub.ai_checks = None, []
            log(session, "reopened", actor, project.id, uid, f.item_id, f"new evidence {code} via smart import")
            stats["approved_attached"] += 1
        touched.add(sub.id)
        log(session, "smart import", actor, project.id, uid, f.item_id, f"{code} {f.filename} ({f.confidence:.0%} · {f.rule[:80]})")
    if batch.table and batch.table.get("grid"):
        t_stats = apply_table(session, project, batch.table, default_unit, actor, touched)
        stats.update({f"table_{k}": v for k, v in t_stats.items()})
    session.flush()
    session.refresh(project)
    workflow.sync(session, project)
    subs = {s.id: s for s in project.submissions}
    for sid in touched:
        s = subs.get(sid)
        if s is None:
            continue
        if auto_submit and s.status in ("pending", "rejected") and not catalogue.item(s.item_id).get("system"):
            workflow.submit(session, s, actor + " (smart import)")
            stats["submitted"] += 1
        if auto_ai and s.status == "submitted":
            ev = [e for e in project.evidence if e.item_id == s.item_id and e.unit_id in (s.unit_id, None)]
            ai_review.apply(s, project, s.unit, ev)
            stats["ai"] += 1
    batch.status, batch.imported_at, batch.stats = "imported", now(), stats
    shutil.rmtree(STAGING / str(batch.id), ignore_errors=True)
    log(session, "smart import completed", actor, project.id,
        comment=f"{stats['files']} files · {stats.get('table_rows', 0)} table rows · {stats['submitted']} submitted · {stats['ai']} AI-reviewed")
    return stats


def discard(session: Session, batch: ImportBatch) -> None:
    batch.status = "discarded"
    shutil.rmtree(STAGING / str(batch.id), ignore_errors=True)
