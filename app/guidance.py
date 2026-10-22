"""Dynamic evidence guidance — what the client must provide, by when, and what is still missing.

For one catalogue item on one packaging unit it combines:

* the legal level of the item (seeds/guidance/levels.json): required by the PPWR text,
  supporting evidence (one accepted way to show a legal requirement), or good practice;
* the article's applicability and application date for *this* project (applicability matrix)
  → "legally required now", "legally required from <date>", or "not required";
* the minimum evidence checklist, evaluated live against the entry, the uploaded files, the BOM
  and the supplier library;
* the AI pre-review rules run as a dry run (no state change), plus field validation;
* contextual hints: deadlines, profile consistency, cross-item consistency, regulatory updates.

Nothing here invents a requirement: every 'required' line cites the provision it comes from;
anything not in the regulation is labelled recommended / good practice.
"""
from __future__ import annotations

import datetime as dt
import io
import json
import re
from dataclasses import dataclass, field
from functools import lru_cache
from types import SimpleNamespace

from . import ai_review, applicability, catalogue, config, rag

GENERAL_APPLICATION = dt.date(2026, 8, 12)
_MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct",
                                       "nov", "dec"], 1)}


@lru_cache(maxsize=1)
def data() -> dict:
    out = {"item_levels": {}, "applies_from": {}, "evidence_types": {}, "items": {}}
    for f in sorted((config.SEEDS_DIR / "guidance").glob("*.json")):
        d = json.loads(f.read_text(encoding="utf-8"))
        for k in out:
            out[k].update(d.get(k, {}))
    return out


def parse_date(text: str) -> dt.date | None:
    if not text:
        return None
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", text)
    if m:
        return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.search(r"(\d{1,2})\s+([A-Za-z]{3})[a-z]*\s+(\d{4})", text)
    if m and m.group(2).lower() in _MONTHS:
        return dt.date(int(m.group(3)), _MONTHS[m.group(2).lower()], int(m.group(1)))
    return None


def level_of(item_id: str) -> tuple[str, str, str]:
    v = data()["item_levels"].get(item_id)
    return tuple(v) if v else ("supporting", "", "")


def evidence_profile(code: str) -> dict:
    return data()["evidence_types"].get(code) or data()["evidence_types"]["DOC"]


def item_guidance(item_id: str) -> dict:
    """Detailed guidance, or a fallback built from the catalogue and the evidence-type profile."""
    it = catalogue.item(item_id)
    g = dict(data()["items"].get(item_id, {}))
    prof = evidence_profile(it["evidence_code"])
    g["detailed"] = item_id in data()["items"]
    g.setdefault("summary", [f"{it['requirement']}. {it['logic']}", f"{it.get('requirement_zh', '')}。{it.get('logic_zh', '')}"])
    g.setdefault("terms", [])
    g.setdefault("logic", [[f"{it['legal']} — {it['logic']}", it.get("logic_zh", "")]])
    g.setdefault("info", prof["content"])
    if "minimum" not in g:
        mins = []
        keys = [f["k"] for f in it.get("fields", [])]
        if keys:
            mins.append(["All fields completed", "所有字段已填写", "fields:" + ",".join(keys)])
        if it.get("table"):
            mins.append(["At least one table row", "表格至少一行", "table"])
        if not it.get("system") and item_id not in ai_review.FIELD_ONLY:
            mins.append([f"{prof['name'][0]} uploaded", f"已上传{prof['name'][1]}", f"file:{it['evidence_code']}"])
        g["minimum"] = mins
    g.setdefault("recommended", [])
    if "steps" not in g:
        g["steps"] = [[f"Obtain the {prof['name'][0].lower()} for this unit ({it['requirement']}).",
                       f"获取该单元的{prof['name'][1]}（{it.get('requirement_zh') or it['requirement']}）。"],
                      ["Check it against the acceptance criteria below.", "对照下方验收标准检查。"]]
        if it.get("fields"):
            g["steps"].append(["Fill in the fields of 'Your entry'.", "填写“您的填写”中的字段。"])
        g["steps"].append([f"Upload it as evidence type {it['evidence_code']}, then run 'Check my entry'.",
                           f"以证据类型 {it['evidence_code']} 上传，然后点击“检查我的填写”。"])
    g.setdefault("good", [[it["example"], ""]] if it.get("example") else [])
    g.setdefault("bad", prof["mistakes"])
    g.setdefault("fields", {})
    g["acceptance"] = prof["acceptance"]
    g["formats"] = prof["formats"]
    g["profile"] = prof
    if not g.get("kb"):
        g["kb"] = [e["id"] for e in rag.for_item(item_id, it["article"], limit=5)]
    return g


# ---------------------------------------------------------------------------
# Obligation & timeline
# ---------------------------------------------------------------------------

@dataclass
class Obligation:
    code: str                 # mandatory_now | mandatory_from | supporting | practice | not_required
    level: str                # required | supporting | practice
    why: str
    why_zh: str
    article: str
    art_status: str
    art_basis: str
    applies_from: dt.date | None
    due: dt.date | None = None
    overdue: bool = False
    days_left: int | None = None
    notes: list = field(default_factory=list)   # [en, zh]

    @property
    def is_legal(self) -> bool:
        return self.code in ("mandatory_now", "mandatory_from")


def obligation(project, unit, item_id: str, done: bool = False) -> Obligation:
    it = catalogue.item(item_id)
    level, why, why_zh = level_of(item_id)
    override = data()["applies_from"].get(item_id)
    art = override[1] if override else it["article"]
    status, basis = applicability.status_for(project, art)
    row = next((r for r in applicability.matrix(project) if r.key == art), None)
    applies = parse_date(override[0]) if override else (parse_date(row.applies_from) if row else GENERAL_APPLICATION)
    if art == "core" and not override:
        applies = GENERAL_APPLICATION
    notes = []
    if item_id == "C09-35":
        levels = {unit.level} if unit else {u.level for u in project.units}
        if levels & {"Grouped", "Transport", "E-commerce"}:
            applies = dt.date(2030, 1, 1)
            notes.append(["Grouped / transport / e-commerce: empty-space ratio ≤ 50 % (Art 24(1)).",
                          "组合 / 运输 / 电商包装：空隙率 ≤ 50%（Art 24(1)）。"])
        else:
            applies = dt.date(2028, 2, 12)
            notes.append(["Sales packaging: empty space reduced to the minimum necessary (Art 24(4)).",
                          "销售包装：空隙须降至必要最低（Art 24(4)）。"])
    today = dt.date.today()
    if status == applicability.NA:
        code = "not_required"
    elif level == "required":
        code = "mandatory_now" if (applies is None or applies <= today) else "mandatory_from"
    else:
        code = level
    due = None
    if code == "mandatory_now":
        due = project.target_date
        notes.append(["Must be in the TD before the packaging is placed on the market (Art 15(2)).",
                      "包装投放市场前须纳入技术文档（Art 15(2)）。"])
    elif code == "mandatory_from":
        due = applies
        notes.append([f"Not yet in force — collect it now so the TD is ready on {applies:%d %b %Y}.",
                      f"尚未生效——建议现在收集，以便技术文档在 {applies:%Y-%m-%d} 前就绪。"])
    elif code in ("supporting", "practice") and status != applicability.NA and applies and applies > today:
        notes.append([f"The related requirement applies from {applies:%d %b %Y}.",
                      f"相关要求自 {applies:%Y-%m-%d} 起适用。"])
    days = (due - today).days if due else None
    return Obligation(code, level, why, why_zh, art, status, basis, applies, due,
                      bool(due and days is not None and days < 0 and not done and code == "mandatory_now"), days, notes)


# ---------------------------------------------------------------------------
# Evaluation
# ---------------------------------------------------------------------------

def evidence_for(project, sub) -> list:
    return [e for e in project.evidence if e.item_id == sub.item_id and e.unit_id in (sub.unit_id, None)]


def _num(v) -> float | None:
    m = re.search(r"(\d+(?:[.,]\d+)?)", str(v or ""))
    return float(m.group(1).replace(",", ".")) if m else None


def _components(project, unit):
    units = [unit] if unit else project.units
    return [c for u in units for c in u.components]


def _check(spec: str, ctx: dict):
    """True = satisfied, False = missing, None = cannot be verified automatically (expert confirms)."""
    d, ev, g = ctx["data"], ctx["evidence"], ctx["g"]
    parts = spec.split(":")
    kind = parts[0]
    filled = lambda k: bool(str(d.get(k, "")).strip())  # noqa: E731
    if kind == "file":
        types = parts[1].split(",") if len(parts) > 1 and parts[1] else None
        return any(types is None or e.type_code in types for e in ev) or (bool(ev) and types is not None and None)
    if kind == "field":
        return filled(parts[1])
    if kind == "fields":
        return all(filled(k) for k in parts[1].split(","))
    if kind == "pattern":
        k = parts[1]
        if not filled(k):
            return False
        pat = (g.get("fields", {}).get(k) or {}).get("pattern")
        return bool(re.search(pat, str(d.get(k)))) if pat else True
    if kind == "minlen":
        return len(str(d.get(parts[1], "")).strip()) >= int(parts[2])
    if kind == "keywords":
        text = str(d.get(parts[1], "")).lower()
        hits = sum(1 for kw in parts[2].split("|") if kw.lower() in text)
        return hits >= int(parts[3])
    if kind in ("num_le", "num_gt"):
        v = _num(d.get(parts[1]))
        if v is None:
            return False
        return v <= float(parts[2]) if kind == "num_le" else v > float(parts[2])
    if kind == "date_within":
        dd = parse_date(str(d.get(parts[1], "")))
        if dd is None:
            dd = max((e.issued_date for e in ev if e.issued_date), default=None)
        if dd is None:
            return None
        return (dt.date.today() - dd).days / 30.44 <= int(parts[2])
    if kind == "bom":
        comps = _components(ctx["project"], ctx["unit"])
        return bool(comps) and all(c.material and c.mass_g for c in comps)
    if kind == "table":
        return any(r and str(r[0]).strip() for r in (ctx["table"] or []))
    if kind == "certs":
        doc_type = parts[1]
        comps = [c for c in _components(ctx["project"], ctx["unit"]) if c.supplier_material is not None]
        today = dt.date.today()
        if not comps:
            return True if ev else None
        missing = [c for c in comps if not any(
            x.doc_type == doc_type and x.status == "valid" and not (x.expires and x.expires < today)
            for x in c.supplier_material.certificates)]
        return True if not missing else (None if ev else False)
    if kind == "pfas":
        app = str(d.get("app", ""))
        if not app:
            return False
        return filled("note") if app.startswith("Not applicable") else bool(ev)
    if kind == "pcr":
        app = str(d.get("app", ""))
        if not app:
            return False
        return filled("pcr") if app.startswith("Applicable") else True
    if kind == "rule":
        oks = [c["ok"] for c in ctx["ai_checks"]]
        return False if False in oks else (None if None in oks else True)
    return None


def validate_fields(item: dict, g: dict, d: dict) -> list[dict]:
    out = []
    for f in item.get("fields", []):
        rule = g.get("fields", {}).get(f["k"]) or {}
        v = str(d.get(f["k"], "")).strip()
        if v and rule.get("pattern") and not re.search(rule["pattern"], v):
            msg = rule.get("msg") or ["Check the format.", "请检查格式。"]
            out.append({"field": f["k"], "label": f["label"], "label_zh": f.get("label_zh", ""),
                        "en": msg[0], "zh": msg[1] if len(msg) > 1 else ""})
    return out


def evaluate(project, sub, unit=None, evidence=None, data_override: dict | None = None,
             table_override: list | None = None) -> dict:
    """Full live guidance state for one submission. Pure: never changes the submission."""
    it = catalogue.item(sub.item_id)
    g = item_guidance(sub.item_id)
    unit = unit if unit is not None else sub.unit
    evidence = evidence if evidence is not None else evidence_for(project, sub)
    d = dict(sub.data or {})
    if data_override:
        d.update(data_override)
    table = table_override if table_override is not None else (sub.table or [])
    ob = obligation(project, unit, sub.item_id, done=sub.status in ("approved", "na"))
    preview = SimpleNamespace(item_id=sub.item_id, data=d, table=table, na_requested=False, na_reason="",
                              status=sub.status)
    ai_checks = [] if it.get("system") else ai_review.review(preview, project, unit, evidence)[1]
    ctx = {"data": d, "evidence": evidence, "g": g, "project": project, "unit": unit, "table": table,
           "ai_checks": ai_checks}
    minimum = []
    for en, zh, spec in g.get("minimum", []):
        minimum.append({"en": en, "zh": zh, "spec": spec, "ok": _check(spec, ctx)})
    field_issues = validate_fields(it, g, d)
    gaps = []
    if ob.code != "not_required":
        for m in minimum:
            if m["ok"] is False:
                gaps.append({"kind": "missing", "en": m["en"], "zh": m["zh"]})
        has_file_rule = any(m["spec"].startswith("file") for m in minimum)
        for c in ai_checks:
            generic = re.match(r"^“.+” completed$", c["desc"]) or (c["desc"] == "Evidence file attached" and has_file_rule)
            if g["detailed"] and generic:
                continue  # already covered by a more specific minimum-evidence line
            if c["ok"] is False and not any(c["desc"] == x.get("en") for x in gaps):
                gaps.append({"kind": "check", "en": c["desc"] + (f" — {c['msg']}" if c.get("msg") else ""),
                             "zh": c.get("desc_zh", ""), "ref": c.get("ref", "")})
        for fi in field_issues:
            gaps.append({"kind": "format", "en": f"{fi['label']}: {fi['en']}",
                         "zh": f"{fi['label_zh'] or fi['label']}：{fi['zh']}"})
    to_verify = [m for m in minimum if m["ok"] is None] + [c for c in ai_checks if c["ok"] is None]
    done_n = sum(1 for m in minimum if m["ok"])
    pct = round(100 * done_n / len(minimum)) if minimum else (100 if not gaps else 0)
    return {"item": it, "g": g, "ob": ob, "minimum": minimum, "ai_checks": ai_checks, "gaps": gaps,
            "to_verify": to_verify, "field_issues": field_issues, "ready_pct": pct,
            "hints": hints(project, sub, unit, evidence, d, ob, g),
            "kb": [e for e in (rag.pool_entry(x) for x in g.get("kb", [])) if e]}


def hints(project, sub, unit, evidence, d: dict, ob: Obligation, g: dict) -> list[dict]:
    """Contextual, explainable hints. Each hint says what triggered it."""
    out = []
    it = catalogue.item(sub.item_id)

    def add(kind, en, zh):
        out.append({"kind": kind, "en": en, "zh": zh})

    if ob.code == "not_required":
        add("info", f"Not required for this packaging: {ob.art_basis}",
            "本包装不适用该项。如项目概况有误，请在概览页修改。")
        return out
    if ob.overdue:
        add("alert", f"Project target date {ob.due:%d %b %Y} has passed and this legally required item is still open.",
            f"项目目标日期 {ob.due:%Y-%m-%d} 已过，而该法定必需项仍未完成。")
    elif ob.code == "mandatory_now" and ob.days_left is not None and ob.days_left <= 30:
        add("alert", f"{ob.days_left} days left until the project target date — legally required item.",
            f"距项目目标日期还有 {ob.days_left} 天——该项为法定必需。")
    if sub.status == "rejected" and sub.review_comment:
        add("alert", f"The expert asked: {sub.review_comment}", f"专家退回意见：{sub.review_comment}")
    expected = it["evidence_code"]
    if evidence and not any(e.type_code == expected for e in evidence) and sub.item_id not in ai_review.FIELD_ONLY:
        add("tip", f"Files are uploaded, but none is of the expected type {expected} "
                   f"({evidence_profile(expected)['name'][0]}). Check the type when uploading.",
            f"已上传文件，但没有预期类型 {expected}（{evidence_profile(expected)['name'][1]}）。上传时请检查类型。")
    bad_ext = [e.filename for e in evidence if e.filename and e.filename.rsplit(".", 1)[-1].lower() not in g["formats"]]
    if bad_ext:
        add("tip", f"Unusual file format for this item: {', '.join(bad_ext[:3])}. Expected: {', '.join(g['formats'][:6])}.",
            f"该项的文件格式不常见：{', '.join(bad_ext[:3])}。预期格式：{', '.join(g['formats'][:6])}。")
    for e in evidence:
        if e.type_code == "TST" and e.issued_date and (dt.date.today() - e.issued_date).days / 30.44 > config.REPORT_MAX_AGE_MONTHS:
            add("tip", f"{e.code} is older than {config.REPORT_MAX_AGE_MONTHS} months — consider a new test (good practice).",
                f"{e.code} 已超过 {config.REPORT_MAX_AGE_MONTHS} 个月——建议重新检测（良好实践）。")
    # cross-item consistency for drawings / artwork revisions
    if sub.item_id in ("C03-10", "C03-11", "C03-12"):
        revs = {}
        for s in project.submissions:
            if s.unit_id == sub.unit_id and s.item_id in ("C03-10", "C03-11", "C03-12"):
                v = (d if s.id == sub.id else (s.data or {})).get("docno", "")
                m = re.search(r"\b(?:R|REV|V|VER)\.?\s?(\d+)", str(v), re.I)
                if m:
                    revs[s.item_id] = m.group(1)
        if len(set(revs.values())) > 1:
            add("alert", "Revisions differ between drawing, die-cut and artwork: "
                + ", ".join(f"{k} R{v}" for k, v in sorted(revs.items())) + " — confirm they describe the same packaging.",
                "图纸、刀模图与 Artwork 版本不一致：" + "，".join(f"{k} R{v}" for k, v in sorted(revs.items()))
                + "——请确认它们描述的是同一包装。")
    if sub.item_id == "C04-70":
        main = next((s for s in project.submissions if s.item_id == "C04-16" and s.unit_id == sub.unit_id), None)
        a, b = _num((main.data or {}).get("res")) if main else None, _num(d.get("res"))
        if a is not None and b is not None and abs(a - b) > 0.5:
            add("alert", f"Result {b:g} mg/kg differs from C04-16 ({a:g} mg/kg).", f"结果 {b:g} mg/kg 与 C04-16（{a:g} mg/kg）不一致。")
    if sub.item_id in ("C04-17", "C04-71") and project.food_contact and str(d.get("app", "")).startswith("Not applicable") \
            and (unit is None or unit.level == "Sales"):
        add("alert", "The project is food-contact and this is a sales unit — 'Not applicable' needs a strong reason.",
            "项目为食品接触且本单元为销售包装——选择“不适用”需充分理由。")
    if ob.article == "a7":
        plastic = [c for c in _components(project, unit) if c.material_group == "plastic"]
        total = sum(c.mass_g or 0 for c in _components(project, unit))
        share = (sum(c.mass_g or 0 for c in plastic) / total) if total else 0
        if plastic and share < 0.05:
            add("tip", f"Plastic is {share:.1%} of the unit mass — below 5 %, Art 7(5)(b) exempts it. Record this as the N/A reason.",
                f"塑料占单元质量 {share:.1%}——低于 5%，依 Art 7(5)(b) 豁免。请将此记录为 N/A 理由。")
    for upd in _updates_for(sub.item_id):
        add("update", f"Regulatory update {upd['code']}: {upd['title']} — {upd['impact']}",
            f"法规动态 {upd['code']}：{upd['title_zh']}——{upd['impact_zh']}")
    if ob.code == "mandatory_from":
        add("info", f"Legally required from {ob.applies_from:%d %b %Y}; until then this is a transitional record.",
            f"自 {ob.applies_from:%Y-%m-%d} 起为法定必需；在此之前为过渡期记录。")
    return out


def _updates_for(item_id: str) -> list[dict]:
    try:
        from . import monitor
        return monitor.updates_for_item(item_id)
    except Exception:
        return []


# ---------------------------------------------------------------------------
# Project-level gap report
# ---------------------------------------------------------------------------

ORDER = {"mandatory_now": 0, "mandatory_from": 1, "supporting": 2, "practice": 3, "not_required": 9}


def project_gaps(project) -> dict:
    rows, counts = [], {k: 0 for k in ORDER}
    ev_by = {}
    for e in project.evidence:
        ev_by.setdefault(e.item_id, []).append(e)
    for s in project.submissions:
        it = catalogue.item(s.item_id)
        if it is None:
            continue
        done = s.status in ("approved", "na")
        ob = obligation(project, s.unit, s.item_id, done=done)
        counts[ob.code] += 0 if done else 1
        if done or ob.code == "not_required" or it.get("system"):
            continue
        evidence = [e for e in ev_by.get(s.item_id, []) if e.unit_id in (s.unit_id, None)]
        g = item_guidance(s.item_id)
        ctx = {"data": s.data or {}, "evidence": evidence, "g": g, "project": project, "unit": s.unit,
               "table": s.table or [], "ai_checks": s.ai_checks or []}
        mins = [(en, zh, _check(spec, ctx)) for en, zh, spec in g.get("minimum", [])]
        missing = [(en, zh) for en, zh, ok in mins if ok is False]
        rows.append({"sub": s, "item": it, "ob": ob, "missing": missing,
                     "met": sum(1 for *_x, ok in mins if ok), "total": len(mins)})
    rows.sort(key=lambda r: (ORDER[r["ob"].code], r["ob"].due or dt.date.max, r["sub"].item_id,
                             r["sub"].unit.code if r["sub"].unit else ""))
    return {"rows": rows, "counts": counts,
            "overdue": sum(1 for r in rows if r["ob"].overdue)}


# ---------------------------------------------------------------------------
# Templates
# ---------------------------------------------------------------------------

def template_docx(item_id: str, project=None, unit=None) -> bytes:
    """A fill-in template for the item: what to provide, an example, acceptance criteria."""
    from docx import Document
    from docx.shared import Pt

    it = catalogue.item(item_id)
    g = item_guidance(item_id)
    level, why, why_zh = level_of(item_id)
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(f"{it['id']} {it['name']} / {it['name_zh']}", 0)
    p = doc.add_paragraph()
    p.add_run("Template — PPWR (EU) 2025/40 · 模板").italic = True
    if project:
        doc.add_paragraph(f"Project 项目: {project.code} {project.title}"
                          + (f" · Unit 单元: {unit.code} {unit.name}" if unit else ""))
    doc.add_heading("Legal basis · 法规依据", 1)
    doc.add_paragraph(it["legal"])
    lvl = {"required": "Required by the regulation (when the article applies) · 法规要求（条款适用时）",
           "supporting": "Recommended evidence · 建议证据", "practice": "Good practice (optional) · 良好实践（可选）"}[level]
    doc.add_paragraph(f"{lvl} — {why} / {why_zh}")
    doc.add_heading("What to provide · 需提供的信息", 1)
    for pair in g.get("info", []):
        doc.add_paragraph(f"{pair[0]} / {pair[1]}", style="List Bullet")
    if it.get("fields"):
        doc.add_heading("Data · 数据", 1)
        tbl = doc.add_table(rows=1, cols=3)
        tbl.style = "Table Grid"
        tbl.rows[0].cells[0].text, tbl.rows[0].cells[1].text, tbl.rows[0].cells[2].text = \
            "Field 字段", "Example 示例", "Your value 填写"
        for f in it["fields"]:
            row = tbl.add_row().cells
            row[0].text = f"{f['label']} / {f.get('label_zh', '')}"
            row[1].text = (g.get("fields", {}).get(f["k"]) or {}).get("ph", "") or " / ".join(f.get("options", []))
    doc.add_heading("Acceptance criteria · 验收标准", 1)
    for pair in g.get("acceptance", []) + [[m[0], m[1]] for m in g.get("minimum", [])]:
        doc.add_paragraph(f"☐ {pair[0]} / {pair[1]}")
    if g.get("good"):
        doc.add_heading("Example of an acceptable submission · 合格示例", 1)
        for pair in g["good"]:
            doc.add_paragraph(" / ".join(x for x in pair if x))
    doc.add_heading("Sign-off · 签署", 1)
    t2 = doc.add_table(rows=3, cols=2)
    t2.style = "Table Grid"
    for i, lab in enumerate(["Prepared by / date 编制人 / 日期", "Approved by / function 批准人 / 职务",
                             "Document no. / revision 文件编号 / 版本"]):
        t2.rows[i].cells[0].text = lab
    doc.add_paragraph().add_run("This template is guidance from the PPWR Digital Platform; it does not add legal "
                                "requirements. 本模板为平台指引，不增加法定要求。").italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
