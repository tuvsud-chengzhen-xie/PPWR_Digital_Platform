"""Excel in/out in the client's own templates (manual collection kit).

- PPWR_BOM_V1.xlsx      one row per component: PS / PU / Component ID, level, material, mass …
- PPWR_Evidence_List_V1 one row per evidence file, Evidence ID PU-XXX-CXXX-EVI-XXX

Clients who already filled the templates by hand can import them; every project
can be exported back in the same layout, so the platform and the manual process
stay interchangeable during roll-out.
"""
from __future__ import annotations

import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from . import catalogue, workflow
from .db import Component, PackagingUnit, Project, Supplier, log

BOM_COLUMNS = [
    ("ps", "Packaging System ID 包装编号"), ("pu", "Packaging Unit ID 包装单元编号"),
    ("comp", "Component ID  包装组件编号"), ("type", "Packaging Type 包装类型"),
    ("level", "Packaging Level 包装层级"), ("desc", "Packaging Description 包装描述"),
    ("structure", "Packaging Structure 包装结构"), ("material", "Material 材质"),
    ("mass", "Component Mass (g) 重量"), ("dims", "Dimension (mm)  尺寸"), ("supplier", "Supplier 供应商"),
    ("pcr", "Recycled Content (%) (if yes)"), ("bio", "Biobased Content (%)  (if yes)"),
    ("articles", "Applicable PPWR Articles"), ("evidence", "Evidence 证明文件"), ("remarks", "Remarks备注"),
]
HEADER_MATCH = {
    "ps": "packaging system", "pu": "packaging unit", "comp": "component id", "type": "packaging type",
    "level": "packaging level", "desc": "description", "structure": "structure", "material": "material",
    "mass": "mass", "dims": "dimension", "supplier": "supplier", "pcr": "recycled", "bio": "biobased",
    "articles": "articles", "evidence": "evidence", "remarks": "remark",
}
LEVELS = {"sales": "Sales", "grouped": "Grouped", "transport": "Transport", "e-commerce": "E-commerce",
          "ecommerce": "E-commerce", "shipper": "Transport", "primary": "Sales", "secondary": "Grouped",
          "tertiary": "Transport"}


def _num(v):
    if v is None or str(v).strip() in ("", "-", "—", "N/A", "n/a"):
        return None
    m = re.search(r"(\d+(?:[.,]\d+)?)", str(v))
    return float(m.group(1).replace(",", ".")) if m else None


def _level(v) -> str:
    t = str(v or "").lower()
    for k, lv in LEVELS.items():
        if k in t:
            return lv
    return "Sales"


def _structure(v) -> str:
    t = str(v or "").lower()
    if "multi" in t:
        return "Multi-layer"
    if "compos" in t:
        return "Composite"
    if "mono" in t:
        return "Mono-material"
    return "Other" if t else "Mono-material"


def import_bom(session: Session, project: Project, data: bytes, actor: str) -> dict:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    header_row, cols = None, {}
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15)):
        texts = [str(c.value or "").lower() for c in r]
        if any("component id" in t for t in texts):
            header_row = r[0].row
            for i, t in enumerate(texts):
                for key, needle in HEADER_MATCH.items():
                    if needle in t and key not in cols:
                        cols[key] = i
                        break
            break
    if header_row is None or "comp" not in cols:
        raise ValueError("No 'Component ID' header found — is this the PPWR BOM template?")

    suppliers = session.query(Supplier).filter(Supplier.client_id == project.client_id).all()
    units = {u.code: u for u in project.units}
    stats = {"units_created": 0, "components_created": 0, "components_updated": 0, "skipped": 0,
             "warnings": []}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        def get(key):
            i = cols.get(key)
            return row[i] if i is not None and i < len(row) else None
        comp_id = str(get("comp") or "").strip()
        if not comp_id or comp_id.lower().startswith("e.g") or not re.match(r"^PU-\d+", comp_id, re.I):
            if any(v not in (None, "") for v in row[:6]) and comp_id and not comp_id.lower().startswith("e.g"):
                stats["skipped"] += 1
            continue
        pu_code = str(get("pu") or comp_id.rsplit("-C", 1)[0]).strip().upper()
        unit = units.get(pu_code)
        if unit is None:
            unit = PackagingUnit(project=project, code=pu_code, name=str(get("desc") or pu_code),
                                 level=_level(get("level")), pack_type=str(get("type") or ""))
            session.add(unit)
            session.flush()
            units[pu_code] = unit
            stats["units_created"] += 1
        comp = next((c for c in unit.components if c.code.upper() == comp_id.upper()), None)
        if comp is None:
            comp = Component(unit=unit, code=comp_id.upper())
            session.add(comp)
            stats["components_created"] += 1
        else:
            stats["components_updated"] += 1
        material = str(get("material") or "")
        comp.pack_type = str(get("type") or comp.pack_type or "")
        comp.description = str(get("desc") or comp.description or "")
        comp.structure = _structure(get("structure"))
        comp.material = material
        comp.material_group = catalogue.material_group_for(material)
        comp.mass_g = _num(get("mass"))
        comp.dims = str(get("dims") or "")
        comp.recycled_pct = _num(get("pcr"))
        comp.biobased_pct = _num(get("bio"))
        comp.remarks = " · ".join(str(x) for x in (get("articles"), get("evidence"), get("remarks")) if x)
        sup_text = str(get("supplier") or "").strip()
        if sup_text:
            match = next((s for s in suppliers if s.name.lower() in sup_text.lower()
                          or sup_text.lower() in s.name.lower()), None)
            if match:
                comp.supplier_id = match.id
                mat = next((m for m in match.materials if m.material_group == comp.material_group), None)
                if mat and comp.supplier_material_id is None:
                    comp.supplier_material_id = mat.id
            else:
                stats["warnings"].append(f"{comp.code}: supplier “{sup_text}” not in the supplier library")
                comp.remarks = (comp.remarks + " · " if comp.remarks else "") + f"Supplier: {sup_text}"
        if comp.mass_g is None:
            stats["warnings"].append(f"{comp.code}: no component mass")
    session.flush()
    for u in units.values():
        if u.mass_g is None and u.components:
            masses = [c.mass_g for c in u.components if c.mass_g]
            u.mass_g = round(sum(masses), 2) if masses else None
    session.flush()
    session.refresh(project)
    workflow.ensure_submissions(session, project)
    log(session, "BOM imported", actor, project.id,
        comment=f"{stats['components_created']} created, {stats['components_updated']} updated")
    return stats


def _sheet(title: str, subtitle: str, headers: list[str]):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws["A1"] = title.replace("_", " ")
    ws["A1"].font = Font(bold=True, size=14, color="003D7A")
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="5A6A7A")
    ws.append([])
    ws.append(["No."] + headers)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="003D7A")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i in range(1, len(headers) + 2):
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = 20 if i > 1 else 6
    return wb, ws


def export_bom(project: Project) -> bytes:
    wb, ws = _sheet("PPWR Packaging BOM",
                    f"{project.client.name} · {project.title} · exported from the PPWR Digital Platform",
                    [h for _, h in BOM_COLUMNS])
    n = 0
    for u in project.units:
        for c in u.components:
            n += 1
            ws.append([n, project.ps_code, u.code, c.code, c.pack_type, u.level, c.description, c.structure,
                       c.material, c.mass_g, c.dims,
                       ", ".join(x for x in ((c.supplier.name if c.supplier else ""),
                                             (c.supplier_material.trace_code if c.supplier_material else "")) if x),
                       c.recycled_pct if c.recycled_pct is not None else "—",
                       c.biobased_pct if c.biobased_pct is not None else "—", "",
                       ", ".join(e.code for e in project.evidence if e.component_id == c.id), c.remarks])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_evidence(project: Project) -> bytes:
    headers = ["Packaging System ID 包装编号", "Packaging Unit ID 包装单元编号", "Component ID  包装组件编号",
               "Packaging Type 包装类型", "Evidence 证明文件", "Evidence Type 文件类型", "Evidence No. 文件号码",
               "Evidence ID 文件编号", "Issuing Institution 发布机构", "Issued Date 发布日期",
               "Responsible 责任主体", "Conformable 符合性", "Supplier 供应商", "Remarks 备注"]
    wb, ws = _sheet("PPWR Evidence List",
                    f"{project.client.name} · {project.title} · exported from the PPWR Digital Platform", headers)
    for n, e in enumerate(project.evidence, 1):
        it = catalogue.item(e.item_id) if e.item_id else None
        ws.append([n, project.ps_code, e.unit.code if e.unit else "", e.component.code if e.component else "",
                   e.component.pack_type if e.component else (e.unit.pack_type if e.unit else ""),
                   e.title, catalogue.EVIDENCE_TYPES.get(e.type_code, e.type_code), e.file_no, e.code,
                   e.issuer, e.issued_date.isoformat() if e.issued_date else "", e.responsible, e.conformity,
                   e.supplier_name, f"{it['id']} {it['name']}" if it else ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
